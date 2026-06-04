"""
Generate all downloadable Office workfiles for the VMware on Azure (AVS)
Advanced Specialization engagement toolkit.

Outputs:
  public/templates/engagement/  – offering, qualification, discovery, WAF, assessment-inputs, DoD
  public/templates/deliverables/ – HLD, LLD, runbook, KT, hypercare
  public/templates/audit/       – evidence-tracker, pre-qual-checklist

Run from repo root:
  python scripts/generate-templates.py
"""

from __future__ import annotations

import os
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt, Inches, RGBColor

from pptx import Presentation
from pptx.util import Inches as PInches, Pt as PPt
from pptx.dml.color import RGBColor as PRGBColor

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
OUT_ENG = ROOT / "public" / "templates" / "engagement"
OUT_DEL = ROOT / "public" / "templates" / "deliverables"
OUT_AUD = ROOT / "public" / "templates" / "audit"
for p in (OUT_ENG, OUT_DEL, OUT_AUD):
    p.mkdir(parents=True, exist_ok=True)

SPEC_NAME = "VMware on Microsoft Azure (Azure VMware Solution)"
SPEC_SHORT = "AVS"
AUTHOR = "Microsoft Partner Innersource Template"
COMPANY = "<Your Partner Org>"

BLUE = RGBColor(0x0F, 0x4C, 0x81)
LIGHT = RGBColor(0xE7, 0xEE, 0xF7)
GREY = RGBColor(0x59, 0x59, 0x59)

PBLUE = PRGBColor(0x0F, 0x4C, 0x81)
PLIGHT = PRGBColor(0xE7, 0xEE, 0xF7)
PGREY = PRGBColor(0x59, 0x59, 0x59)
PWHITE = PRGBColor(0xFF, 0xFF, 0xFF)


# ─── Word helpers ─────────────────────────────────────────────────────────────

def set_core_props(doc: Document, title: str, subject: str = "") -> None:
    cp = doc.core_properties
    cp.title = title
    cp.subject = subject or title
    cp.author = AUTHOR
    cp.company = COMPANY
    cp.keywords = f"Microsoft Advanced Specialization; {SPEC_NAME}"


def add_cover(doc: Document, title: str, subtitle: str) -> None:
    for _ in range(4):
        doc.add_paragraph()
    h = doc.add_paragraph()
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = h.add_run("Microsoft Advanced Specialization")
    r.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = BLUE

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run(SPEC_NAME)
    r.bold = True
    r.font.size = Pt(20)
    r.font.color.rgb = BLUE

    s = doc.add_paragraph()
    s.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = s.add_run(title)
    r.bold = True
    r.font.size = Pt(24)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run(subtitle)
    r.italic = True
    r.font.size = Pt(12)
    r.font.color.rgb = GREY

    for _ in range(2):
        doc.add_paragraph()

    meta = doc.add_table(rows=4, cols=2)
    meta.style = "Light Grid Accent 1"
    for i, (k, v) in enumerate([
        ("Customer", "<customer name>"),
        ("Engagement code", "<code>"),
        ("Version", "0.1 (DRAFT)"),
        ("Date", "<YYYY-MM-DD>"),
    ]):
        meta.rows[i].cells[0].text = k
        meta.rows[i].cells[1].text = v
    doc.add_page_break()


def add_toc(doc: Document) -> None:
    p = doc.add_paragraph()
    r = p.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-3" \\h \\z \\u'
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "Right-click → Update Field to populate the Table of Contents."
    fld_char3 = OxmlElement("w:fldChar")
    fld_char3.set(qn("w:fldCharType"), "end")
    r._r.append(fld_char1)
    r._r.append(instr)
    r._r.append(fld_char2)
    r._r.append(placeholder)
    r._r.append(fld_char3)
    doc.add_page_break()


def add_footer(doc: Document, label: str) -> None:
    for section in doc.sections:
        footer = section.footer
        p = footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(f"{SPEC_NAME} — {label} — Microsoft Partner Innersource Template")
        r.font.size = Pt(8)
        r.font.color.rgb = GREY


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def add_prompt(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.italic = True
    r.font.color.rgb = GREY


def add_para(doc: Document, text: str) -> None:
    doc.add_paragraph(text)


def add_checklist(doc: Document, items: list[str]) -> None:
    for it in items:
        doc.add_paragraph(it, style="List Bullet")


def build_word(name: str, title: str, subtitle: str, sections: list[tuple[str, str]], folder: Path = OUT_ENG) -> Path:
    doc = Document()
    set_core_props(doc, f"Microsoft Advanced Specialization — {SPEC_NAME} — {title}", subtitle)
    add_cover(doc, title, subtitle)
    add_heading(doc, "Table of Contents", level=1)
    add_toc(doc)
    for heading, prompt in sections:
        add_heading(doc, heading, level=1)
        add_prompt(doc, prompt)
        for _ in range(2):
            doc.add_paragraph()
    add_footer(doc, title)
    path = folder / name
    doc.save(str(path))
    return path


# ─── Excel helpers ────────────────────────────────────────────────────────────

HEADER_FILL = PatternFill(start_color="0F4C81", end_color="0F4C81", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
THIN = Side(border_style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def write_header(ws, headers: list[str]) -> None:
    for i, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=i, value=h)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
    ws.freeze_panes = "A2"
    for i, h in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(i)].width = max(14, min(40, len(h) + 4))


def add_status_validation(ws, col_letter: str, last_row: int = 200) -> None:
    dv = DataValidation(type="list", formula1='"Not started,In progress,Blocked,Done"', allow_blank=True)
    dv.add(f"{col_letter}2:{col_letter}{last_row}")
    ws.add_data_validation(dv)


def write_row(ws, row_num: int, values: list) -> None:
    for i, v in enumerate(values, 1):
        c = ws.cell(row=row_num, column=i, value=v)
        c.alignment = Alignment(vertical="top", wrap_text=True)
        c.border = BORDER


def set_xlsx_props(wb: Workbook, title: str) -> None:
    wb.properties.title = f"Microsoft Advanced Specialization — {SPEC_NAME} — {title}"
    wb.properties.creator = AUTHOR
    wb.properties.company = COMPANY
    wb.properties.subject = title
    wb.properties.keywords = f"Microsoft Advanced Specialization; {SPEC_NAME}"


# ─── PowerPoint helpers ───────────────────────────────────────────────────────

def set_pptx_props(prs: Presentation, title: str) -> None:
    cp = prs.core_properties
    cp.title = f"Microsoft Advanced Specialization — {SPEC_NAME} — {title}"
    cp.subject = title
    cp.author = AUTHOR
    cp.company = COMPANY
    cp.keywords = f"Microsoft Advanced Specialization; {SPEC_NAME}"


def add_title_slide(prs: Presentation, title: str, subtitle: str) -> None:
    layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = title
    if len(slide.placeholders) > 1:
        slide.placeholders[1].text = subtitle
    for para in slide.shapes.title.text_frame.paragraphs:
        for run in para.runs:
            run.font.color.rgb = PBLUE
            run.font.bold = True


def add_content_slide(prs: Presentation, title: str, bullets: list[str]) -> None:
    layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = title
    body = slide.placeholders[1]
    tf = body.text_frame
    tf.text = bullets[0] if bullets else ""
    for b in bullets[1:]:
        p = tf.add_paragraph()
        p.text = b
        p.level = 0
    for para in slide.shapes.title.text_frame.paragraphs:
        for run in para.runs:
            run.font.color.rgb = PBLUE
            run.font.bold = True


# ═════════════════════════════════════════════════════════════════════════════
# ENGAGEMENT — Word docs
# ═════════════════════════════════════════════════════════════════════════════

def build_qualification_questionnaire() -> Path:
    sections = [
        ("1. Customer profile", "Capture the customer organisation, industry, regions, regulatory framing, decision-makers, and the named executive sponsor."),
        ("2. Current vSphere estate", "Total clusters, ESXi hosts, vSphere version(s), CPU model(s), vSAN vs traditional storage, vCenter topology, NSX-T or NSX-V status, HCX licensing."),
        ("3. Cluster sizing & utilisation", "vCPU/vRAM/storage per cluster, average CPU/RAM utilisation, peak utilisation, growth trajectory (12-month forward forecast)."),
        ("4. Workload inventory", "Approximate VM count, OS mix, criticality tiers (T0/T1/T2), application owners, licensing constraints (SAP, Oracle, SQL FCI, SQL AGs)."),
        ("5. Networking", "On-prem WAN topology, existing ExpressRoute circuits (location, bandwidth, peering type), planned ER bandwidth uplift, DNS, AD topology, IP plan & non-overlapping CIDR availability for AVS mgmt + workload nets."),
        ("6. Identity & security", "Entra ID tenant, on-prem AD trust, federation, MFA, PIM, Defender for Cloud posture, regulatory controls in scope (HIPAA, PCI, ISO 27001, etc.)."),
        ("7. Backup, DR, BCDR objectives", "Current backup tool, RPO/RTO targets per tier, DR strategy today, target DR strategy on AVS (SRM, HCX, Azure Site Recovery), tabletop frequency."),
        ("8. Migration urgency & drivers", "Datacenter exit date, contract end dates, hardware EoL, M&A, capacity ceiling, cost pressure, ranked top-3 drivers."),
        ("9. Operations model", "Run-state team structure, ITSM tooling, monitoring stack (vROps, SCOM, Azure Monitor), patching cadence, change-management approval flow."),
        ("10. Pre-qualification confirmation", "Confirm Solutions Partner for Infrastructure (Azure) designation; ACR thresholds; ≥1 AVS production customer and ≥2 Module A customers in last 12 months; AZ-104 + AZ-305 holders; (recommended) VCP-DCV / NSX-T / AZ-700."),
        ("11. Audit logistics", "Module B only vs combined Module A+B; preferred audit window; ISSI auditor coordination; checklist version (V1.9.1 active Jan 1 – Jun 30, 2026)."),
        ("12. Next steps", "Decision point: proceed to Discovery Workshop, scope agreement, and signed SOW; named accountable owner per side; target dates."),
    ]
    return build_word("qualification-questionnaire.docx", "Qualification Questionnaire",
                      "Pre-engagement discovery interview", sections)


def build_definition_of_done() -> Path:
    sections = [
        ("1. Engagement-wide DoD", "All audit controls have evidence stored in the evidence tracker; customer sign-off captured; hypercare entry/exit criteria met; lessons-learned issue opened."),
        ("2. Discovery DoD", "Azure Migrate VMware appliance has ≥7 days telemetry; dependency map complete; AVS node sizing options A/B/C with assumptions documented; remediation plan for Not-Ready VMs."),
        ("3. Design DoD", "Approved HLD + LLD; ALZ alignment confirmed; non-overlapping IP plan signed off; ExpressRoute Global Reach plan; NSX-T DFW baseline ruleset; backup + DR design."),
        ("4. Implementation DoD", "AVS private cloud deployed and integrated with hub; HCX (or chosen migration tool) installed and tested with a pilot VM; identity source bound; Content Library populated; monitoring + backup live."),
        ("5. Migration DoD per wave", "Pre-migration checks pass; cutover runbook executed; user acceptance sign-off; post-cutover validation complete; rollback unused or executed cleanly."),
        ("6. Validation DoD", "Functional + performance + DR test results captured; customer-named acceptance email on file."),
        ("7. Post-deployment DoD", "SOPs in place; monitoring + alerting wired; backup tested with restore demonstrated; DR tested with automated failover documented; KT delivered; hypercare plan active."),
        ("8. Audit-readiness DoD", "Evidence tracker shows ✅ for every required control; reviewer assigned and date stamped; pre-qual checklist green; audit slot requested in Partner Center."),
    ]
    return build_word("definition-of-done.docx", "Definition of Done",
                      "Phase-by-phase exit criteria for the AVS engagement", sections)


# ═════════════════════════════════════════════════════════════════════════════
# DELIVERABLES — Word docs
# ═════════════════════════════════════════════════════════════════════════════

def build_hld() -> Path:
    sections = [
        ("1. Executive summary", "Engagement scope, target outcome, business drivers, success measures."),
        ("2. Business context", "Datacenter exit / consolidation / DR / VDI / mission-critical rehost? Top-3 drivers and the AVS reference architecture chosen."),
        ("3. Solution overview", "AVS region(s), node SKU(s) (AV36 / AV36P / AV52 / AV64), cluster count, growth headroom, multi-AZ or stretched cluster posture."),
        ("4. Architecture (logical)", "Logical diagram: AVS private cloud, hub VNet, ER + Global Reach, identity, mgmt, workload tiers, PaaS via Private Endpoint."),
        ("5. Architecture (physical)", "Physical: ER circuit endpoints, AVS mgmt subnets, workload segments, NSX-T tier-0/tier-1 layout, edge cluster."),
        ("6. Connectivity & networking", "ER topology, Global Reach plan, Virtual WAN (if used), DNS forwarder strategy, DHCP placement, MTU."),
        ("7. Identity & access", "Identity source for vCenter/NSX-T, Entra ID integration, on-prem AD trust, RBAC model, PIM."),
        ("8. Storage", "vSAN ESA vs traditional, ANF for AVS use-cases, sizing, snapshots, backup integration."),
        ("9. Migration approach", "HCX vs vMotion-based vs 3rd-party, wave plan summary, pilot VM strategy."),
        ("10. Backup, DR & BCDR", "Azure Backup Server (or partner), SRM/HCX DR, automated DR-runbook approach, RPO/RTO per tier."),
        ("11. Security & governance", "Defender for Cloud, NSX-T DFW baseline, Azure Policy, ALZ alignment."),
        ("12. Operations & monitoring", "Azure Monitor agents on guest VMs, vROps replacement plan if applicable, alerting, runbook tooling, patching."),
        ("13. Cost", "Pricing Calculator output reference, reservations plan, FinOps loop, ANF tiering."),
        ("14. Risks & assumptions", "Open risks, assumptions, dependencies, decisions log."),
        ("15. Approvals", "Customer sign-off block (name, role, date, signature)."),
    ]
    return build_word("hld-template.docx", "High-Level Design (HLD)",
                      "AVS solution design at the architecture level", sections, OUT_DEL)


def build_lld() -> Path:
    sections = [
        ("1. Scope & versions", "AVS portal version, HCX version, NSX-T version, vSphere version, ER circuit IDs, target subscription + RG names."),
        ("2. AVS private cloud build sheet", "Region, AZ, cluster name, node SKU, node count, mgmt CIDR, workload CIDR, HCX-enabled (yes/no), Internet enabled (yes/no)."),
        ("3. Networking — IP plan", "Detailed CIDR plan: hub VNet, AVS mgmt /22, AVS workload segments per app tier, no overlap with on-prem."),
        ("4. Networking — ExpressRoute", "Circuit ID, peering location, SKU, bandwidth, ER Gateway / vWAN Hub, Global Reach pairings, BGP ASNs, MTU."),
        ("5. NSX-T — overlays + DFW", "Tier-0/tier-1 layout, segments per app tier, DFW baseline ruleset (allow mgmt, allow app-to-DB on specific ports, default deny)."),
        ("6. Identity bindings", "vCenter identity source, NSX-T identity source, Entra ID for portal RBAC, on-prem AD DCs on Azure VMs (count, SKU, subnet)."),
        ("7. Storage", "vSAN policy assignment per tier, ANF volume plan (capacity pool, service level, volume size, protocol)."),
        ("8. Backup", "Backup tool, schedule per tier, retention, restore test cadence, MABS/MARS placement."),
        ("9. DR", "SRM/HCX/3rd-party config, protection groups, recovery plans, replication frequency."),
        ("10. Monitoring", "Azure Monitor agents on guests, Log Analytics workspace, custom alerts list, vROps replacement mapping."),
        ("11. Runbook references", "Links to cutover runbook, DR runbook, rollback runbook."),
    ]
    return build_word("lld-template.docx", "Low-Level Design (LLD)",
                      "AVS solution design at the build-sheet level", sections, OUT_DEL)


def build_runbook() -> Path:
    sections = [
        ("1. Wave plan", "List of waves, workload groups per wave, VM counts, migration type per wave (HCX Bulk / vMotion / RAV / Cold), target windows, owners."),
        ("2. Pre-migration checklist", "All pre-flight checks for each wave: HCX service mesh healthy, target segments ready, backup taken, change ticket approved, comms sent."),
        ("3. Cutover steps", "Numbered, time-stamped steps for cutover. Include rollback decision points and named approver."),
        ("4. Post-cutover validation", "Smoke tests, perf checks, monitoring verification, user acceptance prompts."),
        ("5. Rollback plan", "Step-by-step rollback if cutover fails; data divergence handling."),
        ("6. DR runbook (separate exercise)", "Failover steps, failback steps, communication plan, success criteria."),
        ("7. Incident playbook", "Common AVS incidents: HCX service degradation, ER circuit flap, vSAN host failure — triage steps and escalation contacts."),
    ]
    return build_word("runbook-template.docx", "Migration & Operations Runbook",
                      "Wave plan, cutover steps, rollback, DR, incident playbook", sections, OUT_DEL)


def build_kt() -> Path:
    sections = [
        ("1. Audience", "List the customer roles in scope: AVS admin, NSX-T admin, vCenter admin, network team, security team, FinOps team."),
        ("2. Sessions", "Session catalog: AVS portal walk-through, vCenter on AVS, NSX-T basics, HCX operations, Azure-side hub/spoke, backup + DR, FinOps."),
        ("3. Materials", "Slides, recordings, lab guides, environment access details."),
        ("4. Hands-on labs", "Lab list: scale a cluster, add an NSX-T segment, run a test failover, restore a VM, query Azure Monitor logs."),
        ("5. Assessment", "Skills check per role, attestation, owner sign-off."),
        ("6. Reference reading", "Microsoft Learn paths, AVS documentation, internal SOPs."),
    ]
    return build_word("kt-plan-template.docx", "Knowledge Transfer Plan",
                      "Audience, sessions, labs, assessment", sections, OUT_DEL)


def build_hypercare() -> Path:
    sections = [
        ("1. Entry criteria", "Migration complete, customer sign-off captured, all SOPs handed over, monitoring live."),
        ("2. Duration & shifts", "Default 30 calendar days; on-call schedule; business-hours vs after-hours coverage; named primary + secondary on-call."),
        ("3. Severity matrix", "Sev1/Sev2/Sev3/Sev4 definitions and response/resolution SLOs for AVS-specific incidents."),
        ("4. Escalation path", "Partner L1 → Partner L2 → Microsoft AVS support (joint case) → VMware support (when relevant)."),
        ("5. Daily standup format", "Time, attendees, agenda, dashboard reviewed."),
        ("6. Weekly review", "Trend review: incidents, change failures, capacity, cost, customer satisfaction."),
        ("7. Exit criteria", "≤ N Sev2 in last 14 days; SOPs validated end-to-end; customer accepts BAU handover; lessons-learned issue opened."),
        ("8. Handover artifacts", "Updated runbooks, updated SOPs, monitoring/alert tuning history, open-risks register."),
    ]
    return build_word("hypercare-plan-template.docx", "Hypercare Plan",
                      "Entry/exit criteria, on-call, severity, escalation", sections, OUT_DEL)


# ═════════════════════════════════════════════════════════════════════════════
# ENGAGEMENT — PowerPoint
# ═════════════════════════════════════════════════════════════════════════════

def build_offering_pptx() -> Path:
    prs = Presentation()
    prs.slide_width = PInches(13.333)
    prs.slide_height = PInches(7.5)
    set_pptx_props(prs, "Offering One-Pager")
    add_title_slide(prs, "Azure VMware Solution — Migration & Modernisation",
                    f"{SPEC_NAME} • Microsoft Advanced Specialization")
    add_content_slide(prs, "What we deliver", [
        "Datacenter exit, DR, stretched cluster, VDI, hybrid, mission-critical rehost on AVS",
        "Discovery → Design → Migration → Validation → Hypercare",
        "Audit-ready evidence aligned to V1.9.1 controls",
        "Joint Microsoft + VMware support readiness",
    ])
    add_content_slide(prs, "Why this team", [
        "Microsoft Solutions Partner for Infrastructure (Azure)",
        "AZ-104 + AZ-305 certified architects",
        "VMware-skilled delivery practitioners (VCP-DCV / NSX-T recommended)",
        "Proven AVS production references",
    ])
    add_content_slide(prs, "Engagement phases", [
        "Qualification (2 weeks): pre-qual gate, scoping",
        "Discovery (3–4 weeks): Azure Migrate, dependency map, sizing",
        "Design (3 weeks): HLD, LLD, ALZ alignment",
        "Implementation (4–8 weeks): AVS build, HCX, migration waves",
        "Validation + Hypercare (4 weeks): test, sign-off, 30-day support",
    ])
    add_content_slide(prs, "Required Microsoft assessments", [
        "FinOps Review, CASE, Pricing Calculator, DevOps Capability Assessment",
        "Cloud Adoption Security Review, Azure Landing Zone Review",
        "Well-Architected Review (≥ 2 pillars for AVS — Reliability + Cost recommended)",
        "Azure Migrate VMware assessment (performance-based ≥ 7 days)",
    ])
    add_content_slide(prs, "Outcome", [
        "AVS private cloud operational in target region(s)",
        "Workloads migrated with customer-signed acceptance",
        "Audit-ready evidence pack for V1.9.1 (ISSI)",
        "Hypercare exit to BAU support",
    ])
    add_content_slide(prs, "Next steps", [
        "Confirm pre-qualification status",
        "Schedule Discovery Workshop (5 days)",
        "Sign SOW and kick off",
        "Contact: <partner contact name / email>",
    ])
    path = OUT_ENG / "offering-one-pager.pptx"
    prs.save(str(path))
    return path


def build_discovery_pptx() -> Path:
    prs = Presentation()
    prs.slide_width = PInches(13.333)
    prs.slide_height = PInches(7.5)
    set_pptx_props(prs, "Discovery Workshop Deck")
    add_title_slide(prs, "AVS Discovery Workshop", "5-day facilitated workshop")
    add_content_slide(prs, "Agenda", [
        "Day 1 — Business context, drivers, success measures",
        "Day 2 — vSphere estate, workload inventory, dependencies",
        "Day 3 — Networking, identity, security, regulatory",
        "Day 4 — Backup, DR, BCDR, ops model",
        "Day 5 — Findings, sizing options, decision points, next steps",
    ])
    for day, topic, bullets in [
        ("Day 1", "Business context", [
            "Top-3 drivers (DC exit, M&A, EoL, cost, capacity)",
            "Success measures and timeline",
            "Decision-makers and sign-off path",
        ]),
        ("Day 2", "Estate & workloads", [
            "vCenter / cluster inventory",
            "VM count, OS mix, criticality tiers",
            "Application owners and constraints (SAP, Oracle, SQL)",
            "Dependency map from Azure Migrate appliance",
        ]),
        ("Day 3", "Networking, identity, security", [
            "On-prem WAN, ExpressRoute circuits, planned bandwidth",
            "IP plan and non-overlapping CIDR for AVS mgmt / workloads",
            "Entra ID, on-prem AD, federation, MFA, PIM",
            "Regulatory: HIPAA, PCI, ISO 27001, data residency",
        ]),
        ("Day 4", "Backup, DR, ops", [
            "Current backup tool and retention",
            "Target RPO/RTO per tier",
            "DR target: SRM, HCX, Azure Site Recovery, partner",
            "Ops model: ITSM, monitoring, patching, change",
        ]),
        ("Day 5", "Findings & next steps", [
            "AVS sizing options A/B/C with assumptions",
            "Recommended reference architecture",
            "Risks, blockers, decisions required",
            "Proposed wave plan and timeline",
            "SOW outline and approval path",
        ]),
    ]:
        add_content_slide(prs, f"{day} — {topic}", bullets)
    add_content_slide(prs, "Outputs", [
        "Discovery workbook (filled)",
        "Azure Migrate assessment export",
        "Sizing options A/B/C deck",
        "Risk register and decisions log",
        "Proposed SOW",
    ])
    path = OUT_ENG / "discovery-workshop-deck.pptx"
    prs.save(str(path))
    return path


# ═════════════════════════════════════════════════════════════════════════════
# ENGAGEMENT — Excel
# ═════════════════════════════════════════════════════════════════════════════

def build_discovery_workbook() -> Path:
    wb = Workbook()
    set_xlsx_props(wb, "Discovery Workbook")

    sheets = [
        ("Customer", ["Field", "Value"], [
            ("Customer name", ""),
            ("Industry", ""),
            ("Region(s)", ""),
            ("Engagement code", ""),
            ("Executive sponsor", ""),
            ("Partner lead", ""),
            ("Start date", ""),
        ]),
        ("vSphere Estate", ["Cluster", "ESXi hosts", "vSphere ver", "vSAN?", "CPU model", "Avg CPU %", "Peak CPU %", "Total vRAM (GB)", "Notes"], []),
        ("Workloads", ["VM name", "App", "Tier (T0/T1/T2)", "vCPU", "vRAM (GB)", "Disk (GB)", "OS", "Owner", "Migration wave", "Notes"], []),
        ("Networking", ["Item", "Current", "Target", "Notes"], [
            ("On-prem WAN topology", "", "", ""),
            ("ExpressRoute circuit(s)", "", "", ""),
            ("ER bandwidth", "", "", ""),
            ("AVS mgmt CIDR", "n/a", "", "≥ /22 required"),
            ("AVS workload CIDR", "n/a", "", "Non-overlapping"),
            ("DNS forwarder strategy", "", "", ""),
            ("Global Reach pairing", "", "", ""),
        ]),
        ("Identity & Security", ["Item", "Current", "Target", "Notes"], [
            ("Entra ID tenant", "", "", ""),
            ("On-prem AD topology", "", "", ""),
            ("Federation", "", "", ""),
            ("MFA", "", "", ""),
            ("PIM", "", "", ""),
            ("Defender for Cloud posture", "", "", ""),
            ("Regulatory framework(s)", "", "", ""),
        ]),
        ("Backup & DR", ["Item", "Current", "Target", "RPO", "RTO", "Notes"], [
            ("Tier 0 workloads", "", "", "", "", ""),
            ("Tier 1 workloads", "", "", "", "", ""),
            ("Tier 2 workloads", "", "", "", "", ""),
            ("DR strategy", "", "", "", "", ""),
            ("Tabletop frequency", "", "", "", "", ""),
        ]),
        ("Decisions Log", ["Date", "Decision", "Owner", "Status"], []),
        ("Risks", ["Risk", "Impact", "Likelihood", "Mitigation", "Owner", "Status"], []),
    ]

    wb.remove(wb.active)
    for name, headers, rows in sheets:
        ws = wb.create_sheet(title=name)
        write_header(ws, headers)
        for i, row in enumerate(rows, 2):
            write_row(ws, i, list(row))
        for r in range(len(rows) + 2, 30):
            for c in range(1, len(headers) + 1):
                ws.cell(row=r, column=c).border = BORDER
        if "Status" in headers:
            add_status_validation(ws, get_column_letter(headers.index("Status") + 1))

    path = OUT_ENG / "discovery-workshop-workbook.xlsx"
    wb.save(str(path))
    return path


def build_waf_workbook() -> Path:
    wb = Workbook()
    set_xlsx_props(wb, "Well-Architected Review")
    wb.remove(wb.active)

    pillars = {
        "Reliability": [
            "Are SLAs/SLOs/SLIs defined per workload tier on AVS?",
            "Is the deployment multi-AZ or stretched-cluster where needed?",
            "Is automated DR runbook tested at least annually?",
            "Are backup restores tested with evidence captured?",
            "Is host failure tolerance (FTT) policy aligned to tier criticality?",
            "Are health probes and synthetic monitors in place for critical apps?",
            "Is the dependency on ER + Global Reach mapped with failover paths?",
        ],
        "Security": [
            "Is identity source (vCenter, NSX-T) integrated with Entra ID / on-prem AD?",
            "Is RBAC least-privilege applied to AVS Resource Provider and vCenter?",
            "Is Defender for Cloud (Defender for Servers) enrolled on guest VMs?",
            "Is NSX-T DFW baseline ruleset deployed with default-deny?",
            "Are AVS guest VMs patched via Update Mgmt or equivalent?",
            "Is data at rest encrypted (vSAN, ANF, Backup)?",
            "Are Privileged Access Workstations used for AVS admin actions?",
        ],
        "Cost Optimization": [
            "Are AVS node reservations (1y/3y) modelled vs PAYG?",
            "Is rightsizing being applied to migrated VMs?",
            "Is ANF service level appropriate per workload (Standard/Premium/Ultra)?",
            "Is ExpressRoute bandwidth right-sized for steady-state vs migration burst?",
            "Is FinOps Review run with named owner and review cadence?",
            "Are unused VMs/snapshots/orphaned disks reclaimed regularly?",
        ],
        "Operational Excellence": [
            "Are SOPs in place for BAU operations on AVS?",
            "Is monitoring data flowing to Azure Monitor / Log Analytics?",
            "Are change-management approvals enforced via ITSM tool?",
            "Is Infrastructure-as-Code used for AVS-adjacent Azure resources (ALZ)?",
            "Is post-incident review applied with lessons-learned captured?",
            "Is patching cadence formal and tracked?",
        ],
        "Performance Efficiency": [
            "Is sizing validated against performance-based assessment (≥7 days)?",
            "Are noisy-neighbour risks mitigated via DRS rules / affinity?",
            "Is storage performance (vSAN / ANF) validated for top-tier workloads?",
            "Is ER bandwidth/latency validated for migration + steady-state?",
            "Are autoscale / scale-out patterns used where applicable?",
            "Are perf baselines captured before and after migration?",
        ],
    }

    headers = ["Question", "Current State", "Target State", "Gap", "Recommendation", "Owner", "Due Date", "Status"]
    for pillar, qs in pillars.items():
        ws = wb.create_sheet(title=pillar)
        write_header(ws, headers)
        for i, q in enumerate(qs, 2):
            write_row(ws, i, [q, "", "", "", "", "", "", "Not started"])
        add_status_validation(ws, get_column_letter(len(headers)))

    path = OUT_ENG / "waf-assessment.xlsx"
    wb.save(str(path))
    return path


def build_assessment_inputs() -> Path:
    wb = Workbook()
    set_xlsx_props(wb, "Assessment Platform Inputs")
    wb.remove(wb.active)

    sheets = [
        ("Azure Migrate", ["Item", "Value", "Notes"], [
            ("Appliance deployed (yes/no)", "", ""),
            ("Telemetry window (days)", "", "≥ 7 recommended"),
            ("Sizing criterion", "", "Performance-based vs As-on-premises"),
            ("Target AVS region", "", ""),
            ("Target AVS node SKU", "", "AV36 / AV36P / AV52 / AV64"),
            ("Discovery scope (clusters)", "", ""),
            ("Dependency map exported?", "", ""),
            ("Assessment export filename", "", ""),
        ]),
        ("HCX Readiness", ["Item", "Value", "Notes"], [
            ("On-prem HCX Connector deployed", "", ""),
            ("AVS HCX Manager available", "", ""),
            ("Service Mesh state", "", ""),
            ("Bandwidth allocated for migration", "", ""),
            ("Pilot VM migrated successfully", "", ""),
            ("Migration types planned", "", "Bulk / vMotion / RAV / Cold"),
            ("HCX network extension enabled (if applicable)", "", ""),
        ]),
        ("Required Assessments", ["Assessment", "Status", "Customer", "Date", "Owner", "Evidence link"], [
            ("FinOps Review (A.1.1)", "Not started", "", "", "", ""),
            ("CASE (A.1.1)", "Not started", "", "", "", ""),
            ("Pricing Calculator (A.1.2)", "Not started", "", "", "", ""),
            ("DevOps Capability Assessment (A.1.2)", "Not started", "", "", "", ""),
            ("Cloud Adoption Security Review (A.2.1)", "Not started", "", "", "", ""),
            ("Azure Well-Architected Review (A.2.2 + B.2.2)", "Not started", "", "", "", ""),
            ("Azure Landing Zone Review (A.3.1)", "Not started", "", "", "", ""),
            ("Azure Migrate VMware (B.1.1)", "Not started", "", "", "", ""),
            ("Azure Advisor Reliability Review (A.3.3 — optional)", "Not started", "", "", "", ""),
        ]),
    ]

    for name, headers, rows in sheets:
        ws = wb.create_sheet(title=name)
        write_header(ws, headers)
        for i, row in enumerate(rows, 2):
            write_row(ws, i, list(row))
        if "Status" in headers:
            add_status_validation(ws, get_column_letter(headers.index("Status") + 1))

    path = OUT_ENG / "assessment-platform-inputs.xlsx"
    wb.save(str(path))
    return path


# ═════════════════════════════════════════════════════════════════════════════
# AUDIT — Excel
# ═════════════════════════════════════════════════════════════════════════════

CONTROLS = [
    ("A.1.1", "Cloud & AI Adoption Business Strategy", "Module A"),
    ("A.1.2", "Cloud & AI Adoption Plan", "Module A"),
    ("A.2.1", "Security & Governance Tooling", "Module A"),
    ("A.2.2", "Well-Architected Workloads", "Module A"),
    ("A.3.1", "Repeatable Deployment (ALZ)", "Module A"),
    ("A.3.2", "Plan for Skilling", "Module A"),
    ("A.3.3", "Operations Management Tooling", "Module A"),
    ("B.1.1", "Workload Assessment", "Module B"),
    ("B.2.1", "Solution Design", "Module B"),
    ("B.2.2", "WAR of Workloads", "Module B"),
    ("B.3.1", "Infrastructure Implementation & Configuration", "Module B"),
    ("B.3.2", "Migration Tools", "Module B"),
    ("B.3.3", "AVS Configuration & Azure Services Integration", "Module B"),
    ("B.4.1", "Service Validation & Testing", "Module B"),
    ("B.4.2", "Post-Deployment Documentation", "Module B"),
]


def build_evidence_tracker() -> Path:
    wb = Workbook()
    set_xlsx_props(wb, "Evidence Tracker")
    ws = wb.active
    ws.title = "Controls"
    headers = ["Control ID", "Control Title", "Module", "Customer", "Status", "Evidence Link", "Reviewer", "Review Date", "Notes"]
    write_header(ws, headers)
    row = 2
    for cid, title, module in CONTROLS:
        for cust in ("cust1", "cust2", "cust3") if module == "Module A" else ("cust1",):
            write_row(ws, row, [cid, title, module, cust, "Not started", "", "", "", ""])
            row += 1
    add_status_validation(ws, "E", row + 50)

    summary = wb.create_sheet(title="Summary")
    write_header(summary, ["Module", "Total rows", "Done", "% Done"])
    write_row(summary, 2, ["Module A", '=COUNTIF(Controls!C:C,"Module A")', '=COUNTIFS(Controls!C:C,"Module A",Controls!E:E,"Done")', '=IFERROR(C2/B2,0)'])
    write_row(summary, 3, ["Module B", '=COUNTIF(Controls!C:C,"Module B")', '=COUNTIFS(Controls!C:C,"Module B",Controls!E:E,"Done")', '=IFERROR(C3/B3,0)'])
    for r in (2, 3):
        summary.cell(row=r, column=4).number_format = "0%"

    path = OUT_AUD / "evidence-tracker.xlsx"
    wb.save(str(path))
    return path


def build_prequal_checklist() -> Path:
    wb = Workbook()
    set_xlsx_props(wb, "Pre-Qualification Checklist")
    ws = wb.active
    ws.title = "Pre-Qualification"
    headers = ["Area", "Requirement", "Status", "Evidence Link", "Owner", "Notes"]
    write_header(ws, headers)

    rows = [
        ("Solutions Partner", "Active Solutions Partner for Infrastructure (Azure) confirmed in Partner Center", "Not started", "", "", ""),
        ("Solutions Partner", "Screenshot of active designation exported", "Not started", "", "", ""),
        ("ACR", "AVS node consumption threshold confirmed", "Not started", "", "", ""),
        ("ACR", "Adjacent infra (ExpressRoute / ANF / Arc-VMware) captured", "Not started", "", "", ""),
        ("ACR", "PDM verified ACR figure (data may lag 2–4 weeks)", "Not started", "", "", ""),
        ("ACR", "Partner Center ACR export saved", "Not started", "", "", ""),
        ("Customer diversity", "≥ 1 AVS production customer in last 12 months (Module B)", "Not started", "", "", ""),
        ("Customer diversity", "≥ 2 unique customers in last 12 months (Module A)", "Not started", "", "", ""),
        ("Certifications (required)", "AZ-104 (Azure Administrator) — at least 1 holder", "Not started", "", "", ""),
        ("Certifications (required)", "AZ-305 (Azure Solutions Architect Expert) — at least 1 holder", "Not started", "", "", ""),
        ("Certifications (recommended)", "VCP-DCV — VMware delivery skill (not audit gate)", "Not started", "", "", ""),
        ("Certifications (recommended)", "NSX-T — VMware delivery skill (not audit gate)", "Not started", "", "", ""),
        ("Certifications (recommended)", "AZ-700 — Networking depth (not audit gate)", "Not started", "", "", ""),
        ("Module A reuse", "If prior Module A pass < 2 years old, plan to waive Module A at renewal", "Not started", "", "", ""),
        ("Audit logistics", "Audit type chosen: Module B only ($2,400 / 4h) or combined A+B ($3,600 / 8h)", "Not started", "", "", ""),
        ("Audit logistics", "Audit slot requested in Partner Center", "Not started", "", "", ""),
        ("Audit logistics", "Checklist version confirmed: V1.9.1 (active Jan 1 – Jun 30, 2026)", "Not started", "", "", ""),
    ]
    for i, r in enumerate(rows, 2):
        write_row(ws, i, list(r))
    add_status_validation(ws, "C")

    path = OUT_AUD / "pre-qual-checklist.xlsx"
    wb.save(str(path))
    return path


# ═════════════════════════════════════════════════════════════════════════════
# Driver
# ═════════════════════════════════════════════════════════════════════════════

def main() -> None:
    outputs: list[Path] = []
    print("Generating Word documents…")
    outputs.append(build_qualification_questionnaire())
    outputs.append(build_definition_of_done())
    outputs.append(build_hld())
    outputs.append(build_lld())
    outputs.append(build_runbook())
    outputs.append(build_kt())
    outputs.append(build_hypercare())
    print("Generating PowerPoint decks…")
    outputs.append(build_offering_pptx())
    outputs.append(build_discovery_pptx())
    print("Generating Excel workbooks…")
    outputs.append(build_discovery_workbook())
    outputs.append(build_waf_workbook())
    outputs.append(build_assessment_inputs())
    outputs.append(build_evidence_tracker())
    outputs.append(build_prequal_checklist())
    print(f"\nGenerated {len(outputs)} workfiles:")
    for p in outputs:
        rel = p.relative_to(ROOT).as_posix()
        size_kb = p.stat().st_size / 1024
        print(f"  - {rel}  ({size_kb:.1f} KB)")


if __name__ == "__main__":
    main()
